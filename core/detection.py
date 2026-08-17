"""
core.detection
===============
Détecte, pour chaque onglet d'un classeur Excel importé, s'il s'agit du
format source V0 ("Export Familles", large) ou V1 (long), à partir de la
signature en cellule A1 (§2 et §3). L'en-tête n'est pas supposé être
obligatoirement en ligne 1 : certains exports réels comportent une ou
plusieurs lignes vides avant l'en-tête (ex. ASSETS.xlsx / "Export Adherent +
Beneficiaire") ; la détection recherche donc la signature sur les premières
lignes de la feuille avant de conclure à un format inconnu. Effectue en un
seul passage de lecture un état des lieux rapide (nombre de lignes/familles,
présence d'images encodées, taux de CIN manquants) afin d'alimenter les
bandeaux d'alerte de l'interface avant même la conversion complète. Fournit
aussi `lire_lignes_donnees`, le lecteur de lignes utilisé pour la conversion
elle-même : sa position dans la liste renvoyée doit rester alignée sur le
numéro de ligne réel du fichier (§7.1, traçabilité), donc aucune ligne
intermédiaire n'est retirée, même vide.
"""
from __future__ import annotations

from dataclasses import dataclass, field
from io import BytesIO

import openpyxl

from core.mapping import (
    LIEN_ADHERENT, SIGNATURE_V0, SIGNATURE_V1, V0_COL_ADHERENT, V0_COL_CIN,
    V0_COL_IMAGE, V0_COL_NUM_FAMILLE, V1_COL_CIN, V1_COL_LIEN,
    V1_COL_NUM_FAMILLE, VERSION_V0, VERSION_V1, construire_index_entetes,
    normaliser_entete, normaliser_lien,
)

VERSION_INCONNUE = "Inconnue"

# Nombre maximal de lignes examinées en tête de feuille à la recherche de la
# signature d'en-tête (§2/§3). Borné pour rester rapide et éviter qu'une
# ligne de données ne soit prise à tort pour un en-tête sur un fichier sans
# en-tête reconnaissable.
MAX_LIGNES_RECHERCHE_ENTETE = 10


@dataclass
class OngletDetecte:
    """Résultat de la détection pour un onglet donné d'un classeur."""

    nom_fichier: str
    nom_onglet: str
    version: str
    en_tetes: list[str] = field(default_factory=list)
    en_tetes_bruts: list[object] = field(default_factory=list)
    ligne_entete: int = 1
    nb_lignes: int = 0
    nb_familles: int = 0
    image_non_vide: bool = False
    cin_absente: bool = False
    pct_sans_cin: float | None = None
    erreur: str | None = None


def _localiser_entete(feuille) -> tuple[int, list[object]]:
    """Cherche, parmi les `MAX_LIGNES_RECHERCHE_ENTETE` premières lignes de
    la feuille, la première dont la première cellule correspond à une
    signature connue (§2/§3). Renvoie (numéro de ligne 1-based, valeurs de
    cette ligne). Si aucune signature n'est trouvée, retombe sur la ligne 1
    (comportement précédent, pour un format réellement inconnu)."""
    premiere_ligne: list[object] = []
    for i, ligne in enumerate(
        feuille.iter_rows(min_row=1, max_row=MAX_LIGNES_RECHERCHE_ENTETE, values_only=True), start=1
    ):
        valeurs = list(ligne)
        if i == 1:
            premiere_ligne = valeurs
        signature = normaliser_entete(valeurs[0]) if valeurs else ""
        if signature in (SIGNATURE_V0, SIGNATURE_V1):
            return i, valeurs
    return 1, premiere_ligne


def _vide(v: object) -> bool:
    return v is None or str(v).strip() == ""


def lire_lignes_donnees(contenu: bytes, nom_onglet: str, ligne_entete: int = 1) -> list[tuple]:
    """Lit les lignes de données d'un onglet (à partir de la ligne suivant
    `ligne_entete`, qui vaut 1 par défaut). Ne retire QUE les lignes vides en
    toute fin de feuille (artefact fréquent des exports Excel qui déclarent
    une plage utilisée plus grande que le contenu réel) : une ligne vide au
    milieu des données est conservée telle quelle, pour que la position de
    chaque ligne dans la liste renvoyée reste alignée sur son numéro réel
    dans le fichier Excel. `core.convert_v0`/`core.convert_v1` s'appuient sur
    cet alignement (combiné à `ligne_entete`) pour calculer "Ligne source"
    (§7.1) ; le retirer romprait la traçabilité de toutes les lignes
    suivantes.
    """
    classeur = openpyxl.load_workbook(BytesIO(contenu), read_only=True, data_only=True)
    feuille = classeur[nom_onglet]
    lignes = list(feuille.iter_rows(min_row=ligne_entete + 1, values_only=True))

    derniere_non_vide = -1
    for i, ligne in enumerate(lignes):
        if any(c is not None and str(c).strip() != "" for c in ligne):
            derniere_non_vide = i
    return lignes[: derniere_non_vide + 1]


def _analyser_donnees(feuille, version: str, index: dict[str, int], ligne_entete: int) -> tuple[int, int, bool, bool, float | None]:
    """Un seul passage sur les lignes de données (à partir de la ligne
    suivant `ligne_entete`) pour calculer : nombre de lignes non vides,
    nombre de familles, présence d'image non vide (V0), absence de colonne
    CIN, et pourcentage d'adhérents sans CIN."""
    nb_lignes = 0
    familles: set[str] = set()
    image_non_vide = False
    nb_familles_sans_cin = 0

    if version == VERSION_V0:
        i_famille = index.get(V0_COL_NUM_FAMILLE)
        i_image = index.get(V0_COL_IMAGE)
        i_cin = index.get(V0_COL_CIN)
        cin_absente = i_cin is None
        for ligne in feuille.iter_rows(min_row=ligne_entete + 1, values_only=True):
            if not any(c is not None and str(c).strip() != "" for c in ligne):
                continue
            nb_lignes += 1
            if i_famille is not None and i_famille < len(ligne):
                familles.add(str(ligne[i_famille]))
            else:
                familles.add(str(nb_lignes))
            if not image_non_vide and i_image is not None and i_image < len(ligne):
                if not _vide(ligne[i_image]):
                    image_non_vide = True
            if cin_absente or (i_cin is not None and i_cin < len(ligne) and _vide(ligne[i_cin])):
                nb_familles_sans_cin += 1
        nb_familles = len(familles) if familles else nb_lignes
        pct = round(100 * nb_familles_sans_cin / nb_familles, 1) if nb_familles else None
        return nb_lignes, nb_familles, image_non_vide, cin_absente, pct

    # V1 : une ligne = une personne ; on compte les familles distinctes et
    # le taux de CIN manquant côté adhérent uniquement (§5.3).
    i_famille = index.get(V1_COL_NUM_FAMILLE)
    i_lien = index.get(V1_COL_LIEN)
    i_cin = index.get(V1_COL_CIN)
    cin_absente = i_cin is None
    nb_adherents = 0
    for ligne in feuille.iter_rows(min_row=ligne_entete + 1, values_only=True):
        if not any(c is not None and str(c).strip() != "" for c in ligne):
            continue
        nb_lignes += 1
        if i_famille is not None and i_famille < len(ligne):
            familles.add(str(ligne[i_famille]))
        lien_val = ligne[i_lien] if i_lien is not None and i_lien < len(ligne) else None
        if normaliser_lien(lien_val) == LIEN_ADHERENT:
            nb_adherents += 1
            if cin_absente or (i_cin is not None and i_cin < len(ligne) and _vide(ligne[i_cin])):
                nb_familles_sans_cin += 1
    nb_familles = len(familles)
    pct = round(100 * nb_familles_sans_cin / nb_adherents, 1) if nb_adherents else None
    return nb_lignes, nb_familles, False, cin_absente, pct


def detecter_classeur(contenu: bytes, nom_fichier: str) -> list[OngletDetecte]:
    """Ouvre un classeur (bytes) et détecte la version de chacun de ses onglets.

    Un fichier illisible ne doit jamais interrompre le traitement du lot :
    en cas d'erreur d'ouverture, un unique OngletDetecte "en erreur" est
    renvoyé pour que l'appelant puisse l'afficher et continuer avec les
    autres fichiers.
    """
    try:
        classeur = openpyxl.load_workbook(BytesIO(contenu), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 - on veut capturer toute erreur de lecture
        return [
            OngletDetecte(
                nom_fichier=nom_fichier,
                nom_onglet="",
                version=VERSION_INCONNUE,
                erreur=f"Fichier illisible : {exc}",
            )
        ]

    resultats: list[OngletDetecte] = []
    for nom_onglet in classeur.sheetnames:
        try:
            feuille = classeur[nom_onglet]
            ligne_entete, en_tetes_bruts = _localiser_entete(feuille)
            en_tetes = [normaliser_entete(v) for v in en_tetes_bruts]
            signature = en_tetes[0] if en_tetes else ""
            if signature == SIGNATURE_V0:
                version = VERSION_V0
            elif signature == SIGNATURE_V1:
                version = VERSION_V1
            else:
                version = VERSION_INCONNUE

            if version == VERSION_INCONNUE:
                resultats.append(
                    OngletDetecte(
                        nom_fichier=nom_fichier, nom_onglet=nom_onglet, version=version,
                        en_tetes=en_tetes, en_tetes_bruts=en_tetes_bruts, ligne_entete=ligne_entete,
                    )
                )
                continue

            index = construire_index_entetes(en_tetes_bruts)
            nb_lignes, nb_familles, image_non_vide, cin_absente, pct_sans_cin = _analyser_donnees(feuille, version, index, ligne_entete)
            resultats.append(
                OngletDetecte(
                    nom_fichier=nom_fichier, nom_onglet=nom_onglet, version=version,
                    en_tetes=en_tetes, en_tetes_bruts=en_tetes_bruts, ligne_entete=ligne_entete,
                    nb_lignes=nb_lignes, nb_familles=nb_familles,
                    image_non_vide=image_non_vide, cin_absente=cin_absente, pct_sans_cin=pct_sans_cin,
                )
            )
        except Exception as exc:  # noqa: BLE001
            resultats.append(
                OngletDetecte(
                    nom_fichier=nom_fichier,
                    nom_onglet=nom_onglet,
                    version=VERSION_INCONNUE,
                    erreur=f"Onglet illisible : {exc}",
                )
            )
    return resultats
