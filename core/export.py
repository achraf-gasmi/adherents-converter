"""
core.export
=============
Écriture des classeurs Excel du format cible (§1, §7) avec openpyxl : la
feuille "Données" (11 colonnes, formats texte pour préserver les zéros
initiaux, dates dd/mm/yyyy) et la feuille "Anomalies" (§7.1, filtres
automatiques, volets figés). Fournit aussi la construction du rapport global
RAPPORT_ANOMALIES.xlsx avec son onglet de synthèse (§7.3).
"""
from __future__ import annotations

from collections import defaultdict
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from core.anomalies import (
    COLONNES_ANOMALIES, SEVERITE_AVERTISSEMENT, SEVERITE_ERREUR,
    SEVERITE_INFORMATION, Anomalie,
)
from core.mapping import (
    F_CIN, F_CLIENT, F_DATE_AFFILIATION, F_DATE_NAISSANCE, F_FICHIER_SOURCE,
    F_GENDRE, F_LIEN, F_NOM, F_NUM_FAMILLE, F_RANG, F_RIB, TARGET_COLUMNS,
)

FEUILLE_DONNEES = "Données"
FEUILLE_ANOMALIES = "Anomalies"
FEUILLE_SYNTHESE = "Synthèse"

FORMAT_TEXTE = "@"
FORMAT_DATE = "dd/mm/yyyy"

# Index (0-based) des colonnes du format cible nécessitant un format texte,
# pour préserver les zéros initiaux (N° Famille en B et H, RIB, CIN).
_COLONNES_TEXTE = {1, 7, 9, 10}
_COLONNES_DATE = {5, 6}


def _trier_lignes(lignes: list[dict]) -> list[dict]:
    def cle(ligne: dict) -> tuple:
        rang = ligne[F_RANG]
        return (str(ligne[F_NUM_FAMILLE]), rang if rang is not None else 999)

    return sorted(lignes, key=cle)


def _valeur_cellule(ligne: dict, index_colonne: int) -> object:
    champs = [
        F_CLIENT, F_NUM_FAMILLE, F_LIEN, F_GENDRE, F_NOM, F_DATE_NAISSANCE,
        F_DATE_AFFILIATION, F_NUM_FAMILLE, F_RANG, F_RIB, F_CIN,
    ]
    champ = champs[index_colonne]
    valeur = ligne.get(champ)
    return "" if valeur is None else valeur


def _ecrire_feuille_donnees(ws, lignes: list[dict]) -> None:
    ws.append(TARGET_COLUMNS)
    for cellule in ws[1]:
        cellule.font = Font(bold=True)

    for ligne in _trier_lignes(lignes):
        ws.append([_valeur_cellule(ligne, i) for i in range(len(TARGET_COLUMNS))])

    derniere_ligne = ws.max_row
    for indice_col in range(len(TARGET_COLUMNS)):
        lettre = get_column_letter(indice_col + 1)
        if indice_col in _COLONNES_TEXTE:
            for ligne_excel in range(2, derniere_ligne + 1):
                ws[f"{lettre}{ligne_excel}"].number_format = FORMAT_TEXTE
        elif indice_col in _COLONNES_DATE:
            for ligne_excel in range(2, derniere_ligne + 1):
                ws[f"{lettre}{ligne_excel}"].number_format = FORMAT_DATE


def _ecrire_feuille_anomalies(ws, anomalies: list[Anomalie]) -> None:
    ws.append(COLONNES_ANOMALIES)
    for cellule in ws[1]:
        cellule.font = Font(bold=True)

    for anomalie in anomalies:
        ws.append(anomalie.vers_ligne())

    if ws.max_row >= 1:
        derniere_colonne = get_column_letter(len(COLONNES_ANOMALIES))
        ws.auto_filter.ref = f"A1:{derniere_colonne}{ws.max_row}"
    ws.freeze_panes = "A2"


def construire_fichier_cible(lignes: list[dict], anomalies: list[Anomalie]) -> BytesIO:
    """Construit le classeur cible d'un fichier source : feuille "Données"
    (lignes converties) + feuille "Anomalies" (uniquement celles de ce
    fichier, déjà filtrées par l'appelant)."""
    classeur = Workbook()
    ws_donnees = classeur.active
    ws_donnees.title = FEUILLE_DONNEES
    _ecrire_feuille_donnees(ws_donnees, lignes)

    ws_anomalies = classeur.create_sheet(FEUILLE_ANOMALIES)
    _ecrire_feuille_anomalies(ws_anomalies, anomalies)

    tampon = BytesIO()
    classeur.save(tampon)
    tampon.seek(0)
    return tampon


def calculer_resume(lignes: list[dict], anomalies: list[Anomalie]) -> dict[str, dict]:
    """Résumé par fichier source : nombre de lignes exportées, nombre
    d'anomalies par sévérité, et pourcentage de personnes sans anomalie
    bloquante (sévérité Erreur), utilisé par l'onglet Synthèse et par
    l'interface (§7.3)."""
    resume: dict[str, dict] = defaultdict(lambda: {
        "lignes": 0, SEVERITE_ERREUR: 0, SEVERITE_AVERTISSEMENT: 0, SEVERITE_INFORMATION: 0,
        "personnes_en_erreur": set(),
    })

    for ligne in lignes:
        resume[ligne[F_FICHIER_SOURCE]]["lignes"] += 1

    for anomalie in anomalies:
        entree = resume[anomalie.fichier_source]
        entree[anomalie.severite] += 1
        if anomalie.severite == SEVERITE_ERREUR:
            entree["personnes_en_erreur"].add((anomalie.num_famille, anomalie.nom))

    for fichier, entree in resume.items():
        total = entree["lignes"]
        nb_en_erreur = len(entree["personnes_en_erreur"])
        entree["pct_sans_erreur"] = round(100 * (1 - nb_en_erreur / total), 1) if total else 100.0
        del entree["personnes_en_erreur"]

    return dict(resume)


def _ecrire_feuille_synthese(ws, lignes: list[dict], anomalies: list[Anomalie]) -> None:
    resume = calculer_resume(lignes, anomalies)
    entetes = [
        "Fichier source", "Lignes exportées", "Erreurs", "Avertissements",
        "Informations", "Total anomalies", "% lignes sans anomalie bloquante",
    ]
    ws.append(entetes)
    for cellule in ws[1]:
        cellule.font = Font(bold=True)

    total_lignes = total_erreurs = total_avert = total_info = 0
    for fichier in sorted(resume.keys()):
        e = resume[fichier]
        total = e[SEVERITE_ERREUR] + e[SEVERITE_AVERTISSEMENT] + e[SEVERITE_INFORMATION]
        ws.append([fichier, e["lignes"], e[SEVERITE_ERREUR], e[SEVERITE_AVERTISSEMENT], e[SEVERITE_INFORMATION], total, e["pct_sans_erreur"]])
        total_lignes += e["lignes"]
        total_erreurs += e[SEVERITE_ERREUR]
        total_avert += e[SEVERITE_AVERTISSEMENT]
        total_info += e[SEVERITE_INFORMATION]

    ws.append([])
    ligne_totaux = ["TOTAL", total_lignes, total_erreurs, total_avert, total_info, total_erreurs + total_avert + total_info, ""]
    ws.append(ligne_totaux)
    for cellule in ws[ws.max_row]:
        cellule.font = Font(bold=True)


def construire_rapport_anomalies(lignes: list[dict], anomalies: list[Anomalie]) -> BytesIO:
    """Construit RAPPORT_ANOMALIES.xlsx (§7.3) : toutes les anomalies tous
    fichiers confondus + une synthèse fichier x sévérité."""
    classeur = Workbook()
    ws_anomalies = classeur.active
    ws_anomalies.title = FEUILLE_ANOMALIES
    _ecrire_feuille_anomalies(ws_anomalies, anomalies)

    ws_synthese = classeur.create_sheet(FEUILLE_SYNTHESE)
    _ecrire_feuille_synthese(ws_synthese, lignes, anomalies)

    tampon = BytesIO()
    classeur.save(tampon)
    tampon.seek(0)
    return tampon
