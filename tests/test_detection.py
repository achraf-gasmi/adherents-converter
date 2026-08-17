from io import BytesIO

import openpyxl

from core.detection import VERSION_INCONNUE, VERSION_V0, VERSION_V1, detecter_classeur


def _classeur_bytes(entetes, lignes):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(entetes)
    for ligne in lignes:
        ws.append(ligne)
    tampon = BytesIO()
    wb.save(tampon)
    return tampon.getvalue()


def _classeur_bytes_decale(entetes, lignes, ligne_entete):
    """Construit un classeur dont l'en-tête ne commence qu'à `ligne_entete`
    (1-based), les lignes au-dessus étant laissées entièrement vides —
    reproduit un export réel avec une ou plusieurs lignes vides en tête de
    fichier avant le véritable en-tête."""
    wb = openpyxl.Workbook()
    ws = wb.active
    for j, valeur in enumerate(entetes, start=1):
        ws.cell(row=ligne_entete, column=j, value=valeur)
    for i, ligne in enumerate(lignes):
        for j, valeur in enumerate(ligne, start=1):
            ws.cell(row=ligne_entete + 1 + i, column=j, value=valeur)
    tampon = BytesIO()
    wb.save(tampon)
    return tampon.getvalue()


def test_detection_v0():
    entetes = ["MATRICULE_SOCIETE", "COUPLE ASSURE", "Num de famille", "ADHERENT", "DATE_NAISS"]
    contenu = _classeur_bytes(entetes, [["1", None, "F1", "ADHERENT UN", "01/01/1980"]])
    resultats = detecter_classeur(contenu, "TEST.xlsx")
    assert resultats[0].version == VERSION_V0


def test_detection_v1():
    entetes = ["Client", "N° Famille", "Code d'identification", "Rang", "Lien", "Nom"]
    contenu = _classeur_bytes(entetes, [["CLIENT", "F1", 1, 0, "Adhérent", "ADHERENT UN"]])
    resultats = detecter_classeur(contenu, "TEST.xlsx")
    assert resultats[0].version == VERSION_V1


def test_detection_inconnue():
    entetes = ["Colonne A", "Colonne B"]
    contenu = _classeur_bytes(entetes, [["x", "y"]])
    resultats = detecter_classeur(contenu, "TEST.xlsx")
    assert resultats[0].version == VERSION_INCONNUE


def test_fichier_illisible_ne_plante_pas():
    resultats = detecter_classeur(b"ceci n'est pas un fichier xlsx", "CORROMPU.xlsx")
    assert len(resultats) == 1
    assert resultats[0].erreur is not None


def test_ligne_vide_au_milieu_ne_decale_pas_les_numeros_de_ligne_suivants():
    # Bug : une ligne entièrement vide au milieu des données (fréquent dans
    # les exports réels) était retirée avant d'être numérotée, décalant le
    # "Ligne source" de toutes les lignes suivantes -> traçabilité fausse
    # dans le rapport d'anomalies (on ne retrouve plus la bonne cellule en
    # ouvrant le fichier source et en tapant le numéro de ligne indiqué).
    from core.anomalies import CollecteurAnomalies
    from core.convert_v1 import convertir_v1
    from core.detection import lire_lignes_donnees

    entetes = ["Client", "N° Famille", "Code d'identification", "Rang", "Lien", "Nom",
               "Date de naissance", "Gendre", "Date d'affiliation", "RIB", "Identité gouvernementale"]
    ligne_vide = tuple([None] * len(entetes))
    lignes = [
        ["CLIENT", "F1", 1, 0, "Adhérent", "PERSONNE UN", "01/01/1980", "H", "01/01/2020", None, "05705239"],
        list(ligne_vide),  # ligne 3 dans le fichier Excel : entièrement vide
        # ligne 4 : CIN invalide, doit être signalée avec Ligne source = 4
        ["CLIENT", "F2", 1, 0, "Adhérent", "PERSONNE DEUX", "01/01/1980", "H", "01/01/2020", None, "pas-un-cin"],
    ]
    contenu = _classeur_bytes(entetes, lignes)

    lignes_source = lire_lignes_donnees(contenu, "Sheet")
    collecteur = CollecteurAnomalies()
    convertir_v1(lignes_source, entetes, "TEST.xlsx", "Sheet", collecteur)

    anomalie_cin = next(a for a in collecteur.anomalies if a.code == "CIN_INVALIDE")
    assert anomalie_cin.ligne_source == 4


def test_detection_avec_ligne_1_vide_et_entete_en_ligne_2():
    # Cas réel signalé : ASSETS.xlsx / "Export Adherent + Beneficiaire" a une
    # ligne 1 entièrement vide avant le véritable en-tête (ligne 2). La
    # détection ne doit pas se limiter à la ligne 1.
    entetes = ["Client", "N° Famille", "Code d'identification", "Rang", "Lien", "Nom",
               "Date de naissance", "Gendre", "Date d'affiliation", "RIB", "Identité gouvernementale"]
    lignes = [["CLIENT", "F1", 1, 0, "Adhérent", "ADHERENT UN", "01/01/1980", "H", "01/01/2020", None, "05705239"]]
    contenu = _classeur_bytes_decale(entetes, lignes, ligne_entete=2)

    resultats = detecter_classeur(contenu, "ASSETS.xlsx")
    d = resultats[0]
    assert d.version == VERSION_V1
    assert d.ligne_entete == 2
    assert d.en_tetes[:2] == ["Client", "N° Famille"]


def test_conversion_avec_entete_decale_tracabilite_correcte():
    # Bout en bout : une fois l'en-tête localisé en ligne 2, les données
    # doivent être lues à partir de la ligne 3, et le "Ligne source" des
    # anomalies doit pointer vers la vraie ligne Excel (pas ligne-2 comme si
    # l'en-tête était en ligne 1).
    from core.anomalies import CollecteurAnomalies
    from core.convert_v1 import convertir_v1
    from core.detection import lire_lignes_donnees

    entetes = ["Client", "N° Famille", "Code d'identification", "Rang", "Lien", "Nom",
               "Date de naissance", "Gendre", "Date d'affiliation", "RIB", "Identité gouvernementale"]
    lignes = [
        ["CLIENT", "F1", 1, 0, "Adhérent", "PERSONNE UN", "01/01/1980", "H", "01/01/2020", None, "05705239"],
        # ligne 4 dans le fichier Excel (entête=2, +2 lignes de données) : CIN invalide
        ["CLIENT", "F2", 1, 0, "Adhérent", "PERSONNE DEUX", "01/01/1980", "H", "01/01/2020", None, "pas-un-cin"],
    ]
    contenu = _classeur_bytes_decale(entetes, lignes, ligne_entete=2)

    resultats = detecter_classeur(contenu, "ASSETS.xlsx")
    d = resultats[0]
    assert d.version == VERSION_V1

    lignes_source = lire_lignes_donnees(contenu, d.nom_onglet, ligne_entete=d.ligne_entete)
    collecteur = CollecteurAnomalies()
    convertis = convertir_v1(lignes_source, d.en_tetes_bruts, "ASSETS.xlsx", d.nom_onglet, collecteur, ligne_entete=d.ligne_entete)

    assert len(convertis) == 2
    anomalie_cin = next(a for a in collecteur.anomalies if a.code == "CIN_INVALIDE")
    assert anomalie_cin.ligne_source == 4
