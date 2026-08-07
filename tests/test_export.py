from datetime import date
from io import BytesIO

import openpyxl

from core.anomalies import CollecteurAnomalies
from core.export import construire_fichier_cible
from core.mapping import (
    F_CIN, F_CLIENT, F_COLONNE_SOURCE, F_DATE_AFFILIATION, F_DATE_NAISSANCE,
    F_FICHIER_SOURCE, F_GENDRE, F_INDEX_SOURCE, F_LIEN, F_LIGNE_SOURCE,
    F_NOM, F_NUM_FAMILLE, F_ONGLET_SOURCE, F_RANG, F_RANG_SOURCE, F_RIB,
)


def _ligne(num_famille, rang, rib=""):
    return {
        F_FICHIER_SOURCE: "TEST.xlsx", F_ONGLET_SOURCE: "Sheet1", F_LIGNE_SOURCE: 2,
        F_COLONNE_SOURCE: "", F_INDEX_SOURCE: 0, F_RANG_SOURCE: None,
        F_CLIENT: "CLIENT", F_NUM_FAMILLE: num_famille, F_LIEN: "Adhérent", F_GENDRE: "H",
        F_NOM: "TEST PERSONNE", F_DATE_NAISSANCE: date(1980, 1, 1), F_DATE_AFFILIATION: date(2020, 1, 1),
        F_RANG: rang, F_RIB: rib, F_CIN: "05705239",
    }


def test_fichier_exporte_a_deux_feuilles_nommees():
    lignes = [_ligne("F1", 0, rib="05206000051500152949")]
    tampon = construire_fichier_cible(lignes, [])
    classeur = openpyxl.load_workbook(tampon)
    assert classeur.sheetnames == ["Données", "Anomalies"]


def test_colonnes_b_et_h_identiques():
    lignes = [_ligne("F1", 0), _ligne("F2", 0)]
    tampon = construire_fichier_cible(lignes, [])
    classeur = openpyxl.load_workbook(tampon)
    ws = classeur["Données"]
    for ligne_excel in range(2, ws.max_row + 1):
        assert ws.cell(row=ligne_excel, column=2).value == ws.cell(row=ligne_excel, column=8).value


def test_rib_zeros_initiaux_preserves_apres_ecriture():
    lignes = [_ligne("F1", 0, rib="05206000051500152949")]
    tampon = construire_fichier_cible(lignes, [])
    classeur = openpyxl.load_workbook(tampon)
    ws = classeur["Données"]
    valeur_rib = ws.cell(row=2, column=10).value
    assert str(valeur_rib) == "05206000051500152949"
    assert ws.cell(row=2, column=10).number_format == "@"


def test_feuille_anomalies_filtree_par_fichier():
    collecteur = CollecteurAnomalies()
    collecteur.ajouter(fichier_source="A.xlsx", onglet_source="Sheet1", ligne_source=2, code="NOM_NETTOYE", num_famille="F1", nom="X")
    collecteur.ajouter(fichier_source="B.xlsx", onglet_source="Sheet1", ligne_source=2, code="NOM_NETTOYE", num_famille="F2", nom="Y")

    anomalies_a = collecteur.pour_fichier("A.xlsx")
    tampon = construire_fichier_cible([], anomalies_a)
    classeur = openpyxl.load_workbook(tampon)
    ws = classeur["Anomalies"]
    fichiers_presents = {ws.cell(row=r, column=1).value for r in range(2, ws.max_row + 1)}
    assert fichiers_presents == {"A.xlsx"}
